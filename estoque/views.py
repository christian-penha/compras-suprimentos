from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.shortcuts import redirect, render
from django.utils.http import urlencode

from openpyxl import load_workbook

from cadastros.models import Produto
from contas.models import Modulo, Papel
from contas.permissoes import requer_papel
from tenancy.models import Filial

from .models import Almoxarifado, Movimentacao, SaldoEstoque, TipoMovimentacao
from .services import registrar_entrada

SESSAO_PREVIA = "entrada_planilha_previa"


def _parse_planilha(arquivo):
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    ws = wb.active
    ws.reset_dimensions()
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [str(c or "").strip().lower() for c in next(linhas)]

    def coluna(*termos):
        for i, nome in enumerate(cabecalho):
            if any(t in nome for t in termos):
                return i
        return None

    col_codigo = coluna("código", "codigo")
    col_qtd = coluna("quantidade", "estoque", "qtd")
    col_valor = coluna("valor", "custo", "preço", "preco")
    if col_codigo is None or col_qtd is None:
        return None, [f"Colunas de código/quantidade não encontradas. Cabeçalho lido: {cabecalho}"]

    validas, erros = [], []
    for n, linha in enumerate(linhas, start=2):
        codigo = str(linha[col_codigo] or "").strip()
        if not codigo:
            continue
        produto = Produto.objects.filter(codigo=codigo).first()
        if produto is None:
            erros.append(f"Linha {n}: produto '{codigo}' não cadastrado.")
            continue
        try:
            quantidade = Decimal(str(linha[col_qtd] or "0").replace(",", "."))
        except InvalidOperation:
            erros.append(f"Linha {n}: quantidade inválida ('{linha[col_qtd]}') para {codigo}.")
            continue
        if quantidade <= 0:
            continue
        valor = Decimal("0")
        if col_valor is not None and linha[col_valor] not in (None, ""):
            try:
                valor = Decimal(str(linha[col_valor]).replace(",", "."))
            except InvalidOperation:
                erros.append(f"Linha {n}: valor inválido ('{linha[col_valor]}') para {codigo}.")
                continue
        validas.append(
            {"codigo": codigo, "produto": str(produto), "quantidade": str(quantidade), "valor": str(valor)}
        )
    return validas, erros


@requer_papel(Papel.ADMINISTRADOR, modulo=Modulo.ESTOQUE)
def entrada_modelo(request):
    """Planilha de exemplo com as colunas que a entrada espera, já com alguns códigos reais."""
    from django.http import HttpResponse
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Entrada"
    ws.append(["Código", "Produto (só para conferência)", "Quantidade", "Valor unitário"])
    for p in Produto.objects.exclude(codigo__isnull=True).order_by("codigo")[:3]:
        ws.append([p.codigo, p.nome, 10, 5.0])
    for col, largura in zip("ABCD", (18, 44, 14, 16)):
        ws.column_dimensions[col].width = largura
    resposta = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resposta["Content-Disposition"] = 'attachment; filename="modelo-entrada-de-material.xlsx"'
    wb.save(resposta)
    return resposta


@requer_papel(Papel.ADMINISTRADOR, modulo=Modulo.ESTOQUE)
def entrada_planilha(request):
    almoxarifados = Almoxarifado.objects.filter(ativo=True).select_related("filial")

    if request.method == "POST" and request.FILES.get("arquivo"):
        arquivo = request.FILES["arquivo"]
        validas, erros = _parse_planilha(arquivo)
        if validas is None:
            messages.error(request, erros[0])
            return redirect("entrada_planilha")
        request.session[SESSAO_PREVIA] = {
            "linhas": validas,
            "erros": erros,
            "arquivo": arquivo.name,
            "almoxarifado_id": request.POST.get("almoxarifado"),
            "observacao": request.POST.get("observacao", ""),
        }
        return redirect("entrada_planilha")

    if request.method == "POST" and request.POST.get("acao") == "aplicar":
        previa = request.session.pop(SESSAO_PREVIA, None)
        if not previa:
            messages.error(request, "Prévia expirada — envie a planilha novamente.")
            return redirect("entrada_planilha")
        destino = Almoxarifado.objects.get(pk=previa["almoxarifado_id"])
        for linha in previa["linhas"]:
            produto = Produto.objects.get(codigo=linha["codigo"])
            registrar_entrada(
                produto=produto, destino=destino,
                quantidade=Decimal(linha["quantidade"]),
                valor_unitario=Decimal(linha["valor"]),
                responsavel=request.user,
                arquivo_origem=previa["arquivo"],
                observacao=previa["observacao"],
            )
        messages.success(
            request,
            f"{len(previa['linhas'])} entrada(s) registradas em {destino} "
            f"(arquivo {previa['arquivo']}).",
        )
        return redirect("entrada_planilha")

    if request.method == "POST" and request.POST.get("acao") == "cancelar":
        request.session.pop(SESSAO_PREVIA, None)
        return redirect("entrada_planilha")

    previa = request.session.get(SESSAO_PREVIA)
    destino = None
    if previa:
        destino = Almoxarifado.objects.filter(pk=previa["almoxarifado_id"]).first()
    return render(
        request, "estoque/entrada_planilha.html",
        {
            "total_quantidade": sum((Decimal(l["quantidade"]) for l in previa["linhas"]), Decimal("0")) if previa else 0,
            "total_valor": sum((Decimal(l["quantidade"]) * Decimal(l["valor"]) for l in previa["linhas"]), Decimal("0")) if previa else 0,
            "recentes": Movimentacao.objects.filter(tipo=TipoMovimentacao.ENTRADA)
            .select_related("produto", "destino__filial")[:8],
            "almoxarifados": almoxarifados, "previa": previa, "destino": destino,
            "filiais_destino": Filial.objects.filter(pk__in=almoxarifados.values("filial_id")).order_by("nome"),
        },
    )


SITUACOES = [
    ("sem_estoque", "Sem estoque"),
    ("abaixo", "Abaixo do mínimo"),
    ("atencao", "Em atenção"),
    ("ok", "Em dia"),
]


def _situacao(saldo):
    if saldo.quantidade <= 0:
        return "sem_estoque"
    if saldo.estoque_minimo is not None and saldo.quantidade < saldo.estoque_minimo:
        return "abaixo"
    if saldo.ponto_reposicao is not None and saldo.quantidade <= saldo.ponto_reposicao:
        return "atencao"
    return "ok"


@requer_papel(Papel.ADMINISTRADOR, modulo=Modulo.ESTOQUE)
def posicao_estoque(request):
    filial_id = request.GET.get("filial", "")
    almoxarifado_id = request.GET.get("almoxarifado", "")
    busca = request.GET.get("q", "").strip()
    situacao = request.GET.get("situacao", "")

    almoxarifados = Almoxarifado.objects.filter(ativo=True).select_related("filial")
    saldos = SaldoEstoque.objects.filter(almoxarifado__in=almoxarifados).select_related(
        "produto", "almoxarifado__filial"
    )
    if filial_id.isdigit():
        saldos = saldos.filter(almoxarifado__filial_id=filial_id)
        # o almoxarifado tem de ser da unidade escolhida; senão o filtro dele é descartado
        if almoxarifado_id.isdigit() and not almoxarifados.filter(pk=almoxarifado_id, filial_id=filial_id).exists():
            almoxarifado_id = ""
    if almoxarifado_id.isdigit():
        saldos = saldos.filter(almoxarifado_id=almoxarifado_id)
    if busca:
        saldos = saldos.filter(
            Q(produto__nome__icontains=busca)
            | Q(produto__codigo__icontains=busca)
            | Q(produto__especificacao__icontains=busca)
        )

    linhas = []
    for s in saldos.order_by("almoxarifado__filial__sigla", "almoxarifado__nome", "produto__nome"):
        s.situacao = _situacao(s)
        s.valor = s.quantidade * s.custo_medio
        linhas.append(s)

    contagem = {chave: 0 for chave, _ in SITUACOES}
    for s in linhas:
        contagem[s.situacao] += 1
    kpis = {
        "itens": sum(1 for s in linhas if s.quantidade > 0),
        "abaixo": contagem["abaixo"] + contagem["sem_estoque"],
        "atencao": contagem["atencao"],
        "valor": sum((s.valor for s in linhas), Decimal("0")),
    }

    if situacao in contagem:
        linhas = [s for s in linhas if s.situacao == situacao]
    pagina = Paginator(linhas, 50).get_page(request.GET.get("pagina"))

    sem_pagina = {k: v for k, v in request.GET.items() if v and k != "pagina"}
    sem_situacao = {k: v for k, v in sem_pagina.items() if k != "situacao"}
    return render(request, "estoque/posicao.html", {
        "pagina": pagina,
        "total": len(linhas),
        "kpis": kpis,
        "contagem": contagem,
        "situacoes": SITUACOES,
        "situacao": situacao,
        "filiais": Filial.objects.filter(ativo=True).order_by("nome"),
        "almoxarifados": almoxarifados.order_by("filial__sigla", "nome"),
        "filial_id": filial_id,
        "almoxarifado_id": almoxarifado_id,
        "almoxarifado_sel": almoxarifados.filter(pk=almoxarifado_id).first() if almoxarifado_id.isdigit() else None,
        "busca": busca,
        "qs_pagina": urlencode(sem_pagina),
        "qs_situacao": urlencode(sem_situacao),
    })
